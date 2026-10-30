#!/usr/bin/env python
# -*- coding:utf-8 -*-
"""openpyxl 落盘辅助：公式注入防护 + 值收敛 + 工作簿装配（导出单源）。

openpyxl 给单元格赋值时会把「``=`` 开头的字符串」自动判定为公式
（``Cell._bind_value`` → ``data_type = "f"``），保存后 Excel / WPS 打开即按
公式求值。业务数据（用户名、备注、数据集字段值）来自用户输入，直接 ``append``
写盘会形成公式注入通道（外链 / DDE）。本模块提供统一落行入口：写入后把
**字符串**单元格显式标记为文本（``data_type = "s"``），数值 / 时间单元格保持
原类型——防注入的同时不影响数字列的可计算性与单元格格式。

在此之上收敛导出装配样板（此前监控导出 / 报表投递 / 导入失败报告各维护一份
Workbook 组装）：:func:`excel_safe_value` 做 openpyxl 兼容值收敛，
:func:`to_xlsx` / :func:`save_xlsx` 负责多 sheet 装配（表头加粗 + 列宽自适应
+ 全行防注入）。DRF 的流式渲染器（``common/drf/renders/excel.py``）协议不同
（边渲染边写、含表格样式/数据校验），有意不收编。
"""

import datetime as dt_module
import io
import json
import uuid
from collections.abc import Iterable, Mapping
from typing import Any

from django.utils import timezone
from django.utils.functional import Promise


def append_text_safe_row(worksheet: Any, row: Any) -> None:
    """向工作表追加一行，并把字符串单元格强制为文本类型（防公式注入）。

    仅覆盖字符串值：openpyxl 对 ``=`` 开头的字符串会置 ``data_type="f"``，
    显式改回 ``"s"`` 后 Excel 按纯文本展示、不再求值；数值 / 日期 / 布尔
    单元格不受影响。
    """
    worksheet.append(row)
    for cell in worksheet[worksheet.max_row]:
        if isinstance(cell.value, str):
            cell.data_type = "s"


def excel_safe_value(value: Any) -> Any:
    """单元格值收敛（openpyxl 仅支持基础标量，其余类型会抛错或带公式风险）。

    - UUID（FK 主键常见）→ 字符串；
    - 带时区的 datetime/time → 本地 naive（Excel 不接受 tzinfo）；
    - dict/list/tuple → JSON 文本（ensure_ascii=False）；
    - 惰性翻译代理（gettext_lazy）→ 显式 str（openpyxl 只接受真 str 实例）；
    - None / 数值 / 朴素时间 / 布尔 → 原样（保持数字列可计算性）。
    """
    if isinstance(value, uuid.UUID):
        return str(value)
    if isinstance(value, Promise):
        return str(value)
    if isinstance(value, dt_module.datetime) and timezone.is_aware(value):
        return timezone.localtime(value).replace(tzinfo=None)
    if isinstance(value, dt_module.time) and value.tzinfo is not None:
        return value.replace(tzinfo=None)
    if isinstance(value, (dict, list, tuple)):
        return json.dumps(value, ensure_ascii=False, default=str)
    return value


def to_xlsx(
    sheets: Mapping[str, Iterable[Iterable[Any]]],
    *,
    bold_header: bool = True,
    auto_width: bool = True,
) -> bytes:
    """多 sheet 装配为 xlsx 字节（全行经 :func:`append_text_safe_row` 防注入）。

    :param sheets: 有序映射「sheet 标题 → 行序列」，首行为表头（标题截断 31 字符，
        空标题按序号兜底，重名由 openpyxl 自动去重）；
    :param bold_header: 表头行加粗（默认开）；
    :param auto_width: 按列内容自适应列宽（上限 40，默认开）。
    """
    from openpyxl import Workbook
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter

    workbook = Workbook()
    for index, (title, rows) in enumerate(sheets.items()):
        worksheet = workbook.active if index == 0 else workbook.create_sheet()
        worksheet.title = str(excel_safe_value(title))[:31] or f"Sheet{index + 1}"
        materialized = [list(row) for row in rows]
        for row in materialized:
            append_text_safe_row(worksheet, [excel_safe_value(value) for value in row])
        if bold_header and materialized:
            for cell in worksheet[1]:
                cell.font = Font(bold=True)
        if auto_width:
            header = materialized[0] if materialized else []
            for column in range(1, len(header) + 1):
                lengths = [len(str(header[column - 1]))]
                lengths += [len(str(row[column - 1])) for row in materialized[1:] if column - 1 < len(row)]
                worksheet.column_dimensions[get_column_letter(column)].width = min(max(lengths) + 2, 40)
    stream = io.BytesIO()
    workbook.save(stream)
    return stream.getvalue()


def save_xlsx(path: Any, sheets: Mapping[str, Iterable[Iterable[Any]]], **kwargs: Any) -> None:
    """装配并写入文件路径（导入失败报告等落盘场景）。"""
    content = to_xlsx(sheets, **kwargs)
    with open(path, "wb") as fp:
        fp.write(content)
