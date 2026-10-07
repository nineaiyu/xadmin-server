# -*- coding: utf-8 -*-
"""AI 知识库文档管理集成测试。

覆盖：上传（同名覆盖/校验）、列表轻量与详情预览（全文 + 分块摘要）、
检索即时命中、删除与启停的分块联动、仓库文档保护、sync 与上传互不越界、
权限门控。检索面即分块表——上传/删除/停用后无需额外动作即生效。
"""

import pytest

from ai.models.ai import AiKnowledgeChunk, AiKnowledgeDocument
from ai.utils.ai import retrieve

pytestmark = pytest.mark.django_db

KNOWLEDGE_URL = "/api/ai/knowledge-documents"

DOC_CONTENT = """# 内网使用指南

## 登录说明

首次登录使用管理员分配的初始密码，登录后立即修改。

## 报表导出

报表导出在下载中心查看进度，支持 xlsx 格式。
"""


def _upload(client, name="内网使用指南", content=DOC_CONTENT):
    return client.post(KNOWLEDGE_URL, {"name": name, "content": content}, format="json")


class TestUpload:
    def test_upload_creates_document_and_chunks(self, auth_client):
        response = _upload(auth_client)
        assert response.status_code == 200, response.data
        doc = AiKnowledgeDocument.objects.get(title="内网使用指南")
        assert doc.source_type == AiKnowledgeDocument.SourceType.UPLOAD
        assert doc.path == "upload/内网使用指南.md"
        assert doc.chunk_count == 3
        assert AiKnowledgeChunk.objects.filter(source_path=doc.path).count() == 3
        assert doc.creator is not None
        # 检索即时命中（无需额外重建动作）
        hits = retrieve("报表导出在哪里看")
        assert any("下载中心" in item["chunk"].content for item in hits)

    def test_same_name_overwrites_in_place(self, auth_client):
        assert _upload(auth_client).status_code == 200
        updated = _upload(auth_client, content="# 内网使用指南\n\n## 新章节\n\n内容已更新")
        assert updated.status_code == 200
        assert AiKnowledgeDocument.objects.filter(title="内网使用指南").count() == 1
        doc = AiKnowledgeDocument.objects.get(title="内网使用指南")
        assert doc.chunk_count == 2
        assert not AiKnowledgeChunk.objects.filter(source_path=doc.path, content__contains="下载中心").exists()

    @pytest.mark.parametrize(
        "payload",
        [
            {"name": "   ", "content": "内容"},
            {"name": "a/b", "content": "内容"},
            {"name": "a\\b", "content": "内容"},
            {"name": "..", "content": "内容"},
            {"name": "指南", "content": "   "},
        ],
    )
    def test_upload_validation(self, auth_client, payload):
        response = auth_client.post(KNOWLEDGE_URL, payload, format="json")
        assert response.status_code == 400

    def test_upload_content_too_large(self, auth_client):
        response = _upload(auth_client, name="超大文档", content="x" * 200_001)
        assert response.status_code == 400


class TestPreviewAndList:
    def test_retrieve_returns_content_and_chunks(self, auth_client):
        doc_pk = _upload(auth_client).json()["data"]["pk"]
        body = auth_client.get(f"{KNOWLEDGE_URL}/{doc_pk}").json()["data"]
        assert "报表导出在下载中心" in body["content"]
        assert len(body["chunks"]) == 3
        assert body["chunks"][0]["index"] == 0 and body["chunks"][0]["size"] > 0

    def test_chunk_preview_truncated_on_db_side(self, auth_client):
        """分块摘要只展示开头 120 字符：DB 侧裁剪（Left/Length），预览语义与全文长度不变。"""
        long_text = "字" * 300
        _upload(auth_client, name="超长分块文档", content=f"# 超长分块文档\n\n## 长章节\n\n{long_text}")
        doc = AiKnowledgeDocument.objects.get(title="超长分块文档")
        chunk = AiKnowledgeChunk.objects.get(source_path=doc.path, chunk_index=1)
        assert len(chunk.content) > 120
        body = auth_client.get(f"{KNOWLEDGE_URL}/{doc.pk}").json()["data"]
        preview_row = body["chunks"][1]
        assert preview_row["preview"] == chunk.content[:120]
        assert len(preview_row["preview"]) == 120
        assert preview_row["size"] == len(chunk.content)

    def test_list_is_lightweight(self, auth_client):
        _upload(auth_client)
        rows = auth_client.get(KNOWLEDGE_URL).json()["data"]["results"]
        assert rows and "content" not in rows[0] and "chunks" not in rows[0]
        assert rows[0]["chunk_count"] == 3


class TestDeleteAndToggle:
    def test_delete_upload_removes_chunks(self, auth_client):
        doc_pk = _upload(auth_client).json()["data"]["pk"]
        assert auth_client.delete(f"{KNOWLEDGE_URL}/{doc_pk}").status_code == 200
        assert not AiKnowledgeChunk.objects.filter(source_path__startswith="upload/").exists()
        assert retrieve("报表导出在哪里看") == []

    def test_repo_document_delete_rejected(self, auth_client):
        repo_doc = AiKnowledgeDocument.objects.create(
            path="docs/manual.md", title="仓库手册", content="内容", source_type=AiKnowledgeDocument.SourceType.REPO
        )
        response = auth_client.delete(f"{KNOWLEDGE_URL}/{repo_doc.pk}")
        assert response.status_code == 400
        assert AiKnowledgeDocument.objects.filter(pk=repo_doc.pk).exists()

    def test_toggle_active_removes_and_restores_chunks(self, auth_client):
        doc_pk = _upload(auth_client).json()["data"]["pk"]
        off = auth_client.patch(f"{KNOWLEDGE_URL}/{doc_pk}", {"is_active": False}, format="json")
        assert off.status_code == 200, off.data
        assert AiKnowledgeChunk.objects.filter(source_path__startswith="upload/").count() == 0
        assert retrieve("报表导出在哪里看") == []
        on = auth_client.patch(f"{KNOWLEDGE_URL}/{doc_pk}", {"is_active": True}, format="json")
        assert on.status_code == 200, on.data
        assert AiKnowledgeChunk.objects.filter(source_path__startswith="upload/").count() == 3


class TestSyncIsolation:
    def test_sync_keeps_upload_documents(self, tmp_path, monkeypatch, auth_client):
        """关键守护：仓库同步（含清理）不得触碰上传文档及其分块。"""
        from ai.utils import ai as ai_utils
        from ai.utils import ai_knowledge

        _upload(auth_client)
        empty_docs = tmp_path / "docs"
        empty_docs.mkdir()
        # 实现位于 ai_knowledge（ai.utils.ai 仅再导出），patch 目标须与实现同源
        monkeypatch.setattr(ai_knowledge, "DOCS_DIR", empty_docs)
        monkeypatch.setattr(ai_knowledge, "ROOT_DOCS", [])
        summary = ai_utils.sync_knowledge()
        assert summary["removed"] == 0
        assert AiKnowledgeChunk.objects.filter(source_path__startswith="upload/").count() == 3
        assert any("下载中心" in item["chunk"].content for item in retrieve("报表导出在哪里看"))


class TestRepoRebuildAction:
    def test_sync_repo_action(self, auth_client, tmp_path, monkeypatch):
        from ai.utils import ai_knowledge

        docs = tmp_path / "docs"
        docs.mkdir()
        (docs / "guide.md").write_text("# Guide\n\n## Setup\n\n安装说明内容", encoding="utf-8")
        # 实现位于 ai_knowledge（ai.utils.ai 仅再导出），patch 目标须与实现同源
        monkeypatch.setattr(ai_knowledge, "DOCS_DIR", docs)
        monkeypatch.setattr(ai_knowledge, "ROOT_DOCS", [])
        response = auth_client.post(f"{KNOWLEDGE_URL}/sync-repo", {}, format="json")
        assert response.status_code == 200, response.data
        # 异步契约：响应只带任务提交信息，同步摘要改经状态端点轮询（测试档 eager 同步跑完）
        data = response.json()["data"]
        assert data["state"] == "running"
        assert data["status_url"] == "sync-repo/status"
        assert AiKnowledgeDocument.objects.filter(path="docs/guide.md").exists()
        status = auth_client.get(f"{KNOWLEDGE_URL}/sync-repo/status").json()["data"]
        assert status["state"] == "done"
        assert status["summary"]["created"] == 1

    def test_sync_repo_conflict_when_already_running(self, auth_client):
        """单飞：已有同步在跑时返回 1001，不排队、不重复扫盘。"""
        from ai.utils.sync_progress import try_acquire_lock

        assert try_acquire_lock()
        response = auth_client.post(f"{KNOWLEDGE_URL}/sync-repo", {}, format="json")
        assert response.json()["code"] == 1001

    def test_sync_repo_task_error_releases_lock(self, monkeypatch):
        """任务级兜底：同步异常落 error 终态并释放单飞锁（同步入口不会永久卡死）。"""
        from ai.tasks import sync_repo_task
        from ai.utils import ai_knowledge
        from ai.utils.sync_progress import get_status, try_acquire_lock

        def _boom():
            raise RuntimeError("sync exploded")

        monkeypatch.setattr(ai_knowledge, "sync_knowledge", _boom)
        with pytest.raises(RuntimeError):
            sync_repo_task.apply(args=[])
        assert get_status()["state"] == "error"
        assert try_acquire_lock()

    def test_first_poll_after_submit_never_hits_previous_terminal_state(self, auth_client, settings):
        """提交即翻 running：上一轮随通道保留的旧终态从提交时刻起不可命中，
        首轮轮询（任务尚未被 worker 拉起）读到的只能是 running，而非旧摘要。"""
        from django.core.cache import cache

        from ai.utils.sync_progress import SYNC_STATUS_KEY, release_lock

        cache.set(
            SYNC_STATUS_KEY,
            {"state": "done", "summary": {"created": 9, "updated": 8, "removed": 7}, "finished_time": "prev"},
            3600,
        )
        settings.CELERY_TASK_ALWAYS_EAGER = False  # 只验证提交时刻的通道语义，不真跑任务
        response = auth_client.post(f"{KNOWLEDGE_URL}/sync-repo", {}, format="json")
        assert response.status_code == 200, response.data
        status = auth_client.get(f"{KNOWLEDGE_URL}/sync-repo/status").json()["data"]
        assert status["state"] == "running"
        assert "summary" not in status
        assert "finished_time" not in status
        release_lock()


class TestBatchOperations:
    def test_batch_destroy_only_uploads_and_clears_chunks(self, auth_client):
        """批量删除：只删上传文档（repo 静默保留），分块随删除清理。"""
        _upload(auth_client, name="批量甲")
        _upload(auth_client, name="批量乙")
        repo = AiKnowledgeDocument.objects.create(
            path="docs/manual.md",
            title="仓库手册",
            content="内容",
            source_type=AiKnowledgeDocument.SourceType.REPO,
        )
        pks = [
            str(pk)
            for pk in AiKnowledgeDocument.objects.filter(source_type=AiKnowledgeDocument.SourceType.UPLOAD).values_list(
                "pk", flat=True
            )
        ]
        assert len(pks) == 2
        response = auth_client.post(f"{KNOWLEDGE_URL}/batch-destroy", pks, format="json")
        assert response.status_code == 200, response.data
        assert not AiKnowledgeDocument.objects.filter(source_type=AiKnowledgeDocument.SourceType.UPLOAD).exists()
        # 分块必须随删除清理（框架批量删除不触发 perform_destroy，此处逐条清理）
        assert not AiKnowledgeChunk.objects.filter(source_path__startswith="upload/").exists()
        assert AiKnowledgeDocument.objects.filter(pk=repo.pk).exists()

    def test_batch_toggle_updates_chunks(self, auth_client):
        """批量停用 → 分块移除退出检索；批量启用 → 分块重建。"""
        _upload(auth_client, name="批量启停甲")
        _upload(auth_client, name="批量启停乙")
        pks = [
            str(pk)
            for pk in AiKnowledgeDocument.objects.filter(source_type=AiKnowledgeDocument.SourceType.UPLOAD).values_list(
                "pk", flat=True
            )
        ]
        off = auth_client.post(f"{KNOWLEDGE_URL}/batch-toggle", {"pks": pks, "is_active": False}, format="json")
        assert off.status_code == 200, off.data
        assert off.json()["data"]["changed"] == 2
        assert AiKnowledgeChunk.objects.filter(source_path__startswith="upload/").count() == 0
        on = auth_client.post(f"{KNOWLEDGE_URL}/batch-toggle", {"pks": pks, "is_active": True}, format="json")
        assert on.status_code == 200, on.data
        assert on.json()["data"]["changed"] == 2
        assert AiKnowledgeChunk.objects.filter(source_path__startswith="upload/").count() == 6

    def test_batch_toggle_merged_rebuild(self, auth_client, monkeypatch):
        """批量启用走合并重建：多文档只触发一次批量重建，不逐文档循环。"""
        from ai.utils import ai_knowledge

        _upload(auth_client, name="合并启停甲")
        _upload(auth_client, name="合并启停乙")
        pks = [
            str(pk)
            for pk in AiKnowledgeDocument.objects.filter(source_type=AiKnowledgeDocument.SourceType.UPLOAD).values_list(
                "pk", flat=True
            )
        ]
        assert (
            auth_client.post(
                f"{KNOWLEDGE_URL}/batch-toggle", {"pks": pks, "is_active": False}, format="json"
            ).status_code
            == 200
        )
        calls = []
        original = ai_knowledge.rebuild_chunks_bulk

        def _spy(documents):
            calls.append([doc.path for doc in documents])
            return original(documents)

        monkeypatch.setattr(ai_knowledge, "rebuild_chunks_bulk", _spy)
        on = auth_client.post(f"{KNOWLEDGE_URL}/batch-toggle", {"pks": pks, "is_active": True}, format="json")
        assert on.status_code == 200, on.data
        assert on.json()["data"]["changed"] == 2
        assert len(calls) == 1 and len(calls[0]) == 2
        assert AiKnowledgeChunk.objects.filter(source_path__startswith="upload/").count() == 6
        for doc in AiKnowledgeDocument.objects.filter(source_type=AiKnowledgeDocument.SourceType.UPLOAD):
            assert doc.is_active and doc.chunk_count == 3

    @pytest.mark.parametrize("payload", [{"pks": []}, {"is_active": True}])
    def test_batch_toggle_validation(self, auth_client, payload):
        response = auth_client.post(f"{KNOWLEDGE_URL}/batch-toggle", payload, format="json")
        assert response.status_code == 400


class TestPermissions:
    def test_anonymous_rejected(self, api_client):
        assert api_client.get(KNOWLEDGE_URL).status_code == 401

    def test_normal_user_without_menu_rejected(self, normal_user):
        from rest_framework.test import APIClient

        client = APIClient(HTTP_USER_AGENT="pytest-agent")
        client.force_authenticate(user=normal_user)
        assert client.get(KNOWLEDGE_URL).status_code == 403

    def test_normal_user_with_menu_allowed(self, normal_user):
        from rest_framework.test import APIClient

        from system.models import Menu, MenuMeta

        meta = MenuMeta.objects.create(title="list:AiKnowledge")
        menu = Menu.objects.create(
            name="list:AiKnowledge",
            path="api/ai/knowledge-documents$",
            method="GET",
            menu_type=Menu.MenuChoices.PERMISSION,
            meta=meta,
        )
        normal_user.roles.first().menu.set([menu])
        client = APIClient(HTTP_USER_AGENT="pytest-agent")
        client.force_authenticate(user=normal_user)
        assert client.get(KNOWLEDGE_URL).status_code == 200


# ---------------------------------------------------------------- 7.3 二进制文档解析


def _build_pdf(text: str) -> bytes:
    """构造含单行文本的最小合法 PDF（pypdf 可解析；测试专用）。"""
    import io

    stream = f"BT /F1 12 Tf 72 720 Td ({text}) Tj ET".encode("latin-1")
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 4 0 R "
        b"/Resources << /Font << /F1 5 0 R >> >> >>",
        b"<< /Length " + str(len(stream)).encode() + b" >>\nstream\n" + stream + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    out = io.BytesIO()
    out.write(b"%PDF-1.4\n")
    offsets = []
    for index, body in enumerate(objects, start=1):
        offsets.append(out.tell())
        out.write(f"{index} 0 obj\n".encode() + body + b"\nendobj\n")
    xref_pos = out.tell()
    out.write(f"xref\n0 {len(objects) + 1}\n".encode())
    out.write(b"0000000000 65535 f \n")
    for offset in offsets:
        out.write(f"{offset:010d} 00000 n \n".encode())
    out.write(f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref_pos}\n%%EOF\n".encode())
    return out.getvalue()


def _build_docx(text: str) -> bytes:
    """构造含单段文本的最小 DOCX（python-docx 生成；测试专用）。"""
    import io

    import docx

    out = io.BytesIO()
    docx.Document().save(out)  # 空文档建立包结构后追加段落
    document = docx.Document(io.BytesIO(out.getvalue()))
    document.add_paragraph(text)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


class TestBinaryDocumentUpload:
    """PDF/DOCX 上传：解析为纯文本入库，检索/分块链路零改动。"""

    def test_pdf_upload_extracts_text_and_searchable(self, auth_client):
        import base64

        payload = {
            "name": "平台白皮书",
            "file_type": "pdf",
            "file_b64": base64.b64encode(_build_pdf("xadmin white paper chapter one")).decode(),
        }
        response = auth_client.post(KNOWLEDGE_URL, payload, format="json")
        assert response.status_code == 200, response.data
        doc = AiKnowledgeDocument.objects.get(title="平台白皮书")
        assert "white paper" in doc.content
        assert AiKnowledgeChunk.objects.filter(source_path=doc.path).exists()
        hits = retrieve("white paper")
        assert any("white paper" in item["chunk"].content for item in hits)

    def test_docx_upload_extracts_text(self, auth_client):
        import base64

        payload = {
            "name": "入职手册",
            "file_type": "docx",
            "file_b64": base64.b64encode(_build_docx("入职手册正文内容")).decode(),
        }
        response = auth_client.post(KNOWLEDGE_URL, payload, format="json")
        assert response.status_code == 200, response.data
        doc = AiKnowledgeDocument.objects.get(title="入职手册")
        assert "入职手册正文内容" in doc.content

    def test_binary_upload_requires_both_fields(self, auth_client):
        import base64

        response = auth_client.post(
            KNOWLEDGE_URL, {"name": "半载荷", "file_b64": base64.b64encode(_build_pdf("x")).decode()}, format="json"
        )
        assert response.status_code == 400
        response = auth_client.post(KNOWLEDGE_URL, {"name": "半载荷", "file_type": "pdf"}, format="json")
        assert response.status_code == 400
        assert not AiKnowledgeDocument.objects.filter(title="半载荷").exists()

    def test_invalid_base64_rejected(self, auth_client):
        response = auth_client.post(
            KNOWLEDGE_URL, {"name": "坏载荷", "file_type": "pdf", "file_b64": "not-base64!!"}, format="json"
        )
        assert response.status_code == 400

    def test_broken_pdf_rejected(self, auth_client):
        import base64

        response = auth_client.post(
            KNOWLEDGE_URL,
            {"name": "坏PDF", "file_type": "pdf", "file_b64": base64.b64encode(b"%PDF-1.4 broken").decode()},
            format="json",
        )
        assert response.status_code == 400
        assert not AiKnowledgeDocument.objects.filter(title="坏PDF").exists()

    def test_unsupported_binary_type_rejected(self, auth_client):
        import base64

        response = auth_client.post(
            KNOWLEDGE_URL,
            {"name": "exe文档", "file_type": "exe", "file_b64": base64.b64encode(b"MZ").decode()},
            format="json",
        )
        assert response.status_code == 400

    def test_text_upload_contract_unchanged(self, auth_client):
        """文本直传路径零变化（回归护栏）。"""
        assert _upload(auth_client, name="纯文本回归").status_code == 200
