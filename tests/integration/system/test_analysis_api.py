# -*- coding: utf-8 -*-
"""大屏与定时报表集成测试。

覆盖：Screen/Report CRUD 与越权、报表调度命中（daily/weekly/monthly × 时刻）、
执行（创建者权限上下文 fail-closed、xlsx 产物进下载中心、邮件降级）、
run 动作派发契约（ExportRecord.pk == task_id）。
"""

from datetime import datetime

import pytest
from django.core import mail
from django.utils import timezone

from system.models import ModelLabelField, UserInfo
from system.models.dataset import Dashboard, Dataset, Report, Screen

pytestmark = pytest.mark.django_db

SCREEN_URL = "/api/system/screens"
REPORT_URL = "/api/system/reports"


@pytest.fixture
def model_registry(db):
    root, _ = ModelLabelField.objects.get_or_create(
        name="system.userinfo", defaults={"field_type": ModelLabelField.FieldChoices.DATA, "label": "用户"}
    )
    for name in ("username", "nickname"):
        ModelLabelField.objects.get_or_create(
            name=name, parent=root, defaults={"field_type": ModelLabelField.FieldChoices.DATA, "label": name}
        )
    return root


@pytest.fixture
def dataset(model_registry, superuser):
    return Dataset.objects.create(
        name="用户清单",
        bound_model="system.userinfo",
        columns=["username", "nickname"],
        visibility="shared",
        creator=superuser,
    )


@pytest.fixture
def screen_urls():
    from system.models import Menu, MenuMeta

    def _make(name, path, method):
        meta = MenuMeta.objects.create(title=name)
        return Menu.objects.create(
            name=name, path=path, method=method, menu_type=Menu.MenuChoices.PERMISSION, meta=meta
        )

    detail = "api/system/screens/(?P<pk>[^/.]+)"
    report_detail = "api/system/reports/(?P<pk>[^/.]+)"
    menus = [
        _make("list:Screen", "api/system/screens$", "GET"),
        _make("create:Screen", "api/system/screens$", "POST"),
        _make("partialUpdate:Screen", detail + "$", "PATCH"),
        _make("list:Report", "api/system/reports$", "GET"),
        _make("create:Report", "api/system/reports$", "POST"),
        _make("partialUpdate:Report", report_detail + "$", "PATCH"),
        _make("run:Report", report_detail + "/run$", "POST"),
    ]
    return menus


def grant(user, menus):
    from system.models import UserRole

    role = UserRole.objects.create(name=f"role-{user.username}", code=user.username)
    user.roles.add(role)
    role.menu.set(menus)


class TestScreen:
    def test_anonymous_rejected(self, api_client):
        assert api_client.get(SCREEN_URL).status_code == 401

    def test_crud_and_visibility(self, auth_client, normal_user, dashboard_factory, screen_urls):
        dashboard_factory()
        payload = {
            "name": "值班大屏",
            "dashboards": [str(Dashboard.objects.first().pk)],
            "interval": 20,
            "refresh": 60,
            "visibility": "shared",
        }
        response = auth_client.post(SCREEN_URL, payload, format="json")
        assert response.status_code == 200, response.data

        # personal 仅创建者可见；other 无菜单 403，授菜单后可见性生效
        mine = Screen.objects.create(name="私人屏", creator=normal_user)
        other = _second_user(screen_urls)
        visible = {item["name"] for item in other["client"].get(SCREEN_URL).json()["data"]["results"]}
        assert "值班大屏" in visible and "私人屏" not in visible

        # 共享只读
        response = other["client"].patch(
            f"{SCREEN_URL}/{Screen.objects.get(name='值班大屏').pk}", {"name": "改"}, format="json"
        )
        assert response.json()["code"] == 1003
        assert mine.pk  # noqa: B015

    def test_interval_bounds(self, auth_client, dashboard_factory):
        dashboard_factory()
        response = auth_client.post(
            SCREEN_URL,
            {"name": "坏间隔", "dashboards": [str(Dashboard.objects.first().pk)], "interval": 1},
            format="json",
        )
        assert response.status_code == 400

    def test_unknown_dashboard_rejected(self, auth_client):
        response = auth_client.post(
            SCREEN_URL, {"name": "坏引用", "dashboards": ["00000000-0000-0000-0000-000000000000"]}, format="json"
        )
        assert response.status_code == 400


class TestReportSchedule:
    def test_daily_match(self, dataset, superuser):
        report = _make_report(dataset, superuser, frequency="daily", send_time="08:00")
        due = timezone.localtime().replace(hour=8, minute=0)
        from system.analysis_tasks import report_due

        assert report_due(report, due) is True
        assert report_due(report, due.replace(hour=9)) is False

    def test_weekday_match(self, dataset, superuser):
        report = _make_report(dataset, superuser, frequency="weekly", send_time="08:00", weekday=0)
        from system.analysis_tasks import report_due

        monday = timezone.localtime().replace(hour=8, minute=0)
        while monday.weekday() != 0:
            monday = monday + timezone.timedelta(days=1)
        assert report_due(report, monday) is True
        tuesday = monday + timezone.timedelta(days=1)
        assert report_due(report, tuesday) is False

    def test_monthly_first_day(self, dataset, superuser):
        report = _make_report(dataset, superuser, frequency="monthly", send_time="08:00")
        from system.analysis_tasks import report_due

        first = timezone.localtime().replace(day=1, hour=8, minute=0)
        assert report_due(report, first) is True
        assert report_due(report, first.replace(day=2)) is False

    def test_recipients_email_validated(self, auth_client, dataset):
        payload = {"name": "坏邮箱", "dataset": str(dataset.pk), "recipients": ["not-an-email"]}
        response = auth_client.post(REPORT_URL, payload, format="json")
        assert response.status_code == 400

    def test_aggregate_requires_group_by(self, auth_client, dataset):
        payload = {"name": "缺分组", "dataset": str(dataset.pk), "mode": "aggregate", "recipients": ["a@corp.com"]}
        response = auth_client.post(REPORT_URL, payload, format="json")
        assert response.status_code == 400


class TestReportRun:
    def test_run_creates_record_and_email(self, auth_client, dataset, superuser, settings):
        """run 动作：预创建 ExportRecord（下载中心）→ 任务执行 → 邮件附件。"""
        settings.EMAIL_BACKEND = "django.core.mail.backends.locmem.EmailBackend"
        report = _make_report(dataset, superuser, recipients=["boss@corp.com"])
        response = auth_client.post(f"{REPORT_URL}/{report.pk}/run", {}, format="json")
        assert response.status_code == 200, response.data
        task_id = response.json()["data"]["task_id"]

        from system.models.export import ExportRecord

        record = ExportRecord.objects.get(pk=task_id)
        assert record.status == ExportRecord.Status.SUCCESS
        assert record.file and record.file.filesize > 0
        assert len(mail.outbox) == 1
        assert mail.outbox[0].to == ["boss@corp.com"]
        assert mail.outbox[0].attachments
        report.refresh_from_db()
        assert report.last_status == "SUCCESS"

    def test_run_fail_closed_for_creator_without_grant(self, normal_user, dataset, model_registry, screen_urls):
        """核心验收：创建者无数据授权 → 产物为空且 SUCCESS（fail-closed 不报错）。"""
        from openpyxl import load_workbook

        grant(normal_user, screen_urls)
        report = _make_report(dataset, normal_user, recipients=["boss@corp.com"])
        from system.analysis_tasks import schedule_report_run

        task_id = schedule_report_run(str(report.pk))
        from system.models.export import ExportRecord

        record = ExportRecord.objects.get(pk=task_id)
        assert record.status == ExportRecord.Status.SUCCESS
        sheet = load_workbook(record.file.filepath).active
        assert sheet.max_row == 1  # 只有表头：无授权 → 空结果（fail-closed）

    def test_run_renders_fk_uuid_column(self, auth_client, dataset, superuser, model_registry, settings):
        """FK 列（UUID pk）进 xlsx：单元格转字符串，不再抛 Cannot convert UUID（回归守护）。"""
        from openpyxl import load_workbook

        from system.models import DeptInfo
        from system.models.export import ExportRecord

        settings.EMAIL_BACKEND = "django.core.mail.backends.locmem.EmailBackend"
        root = ModelLabelField.objects.get(name="system.userinfo")
        ModelLabelField.objects.get_or_create(
            name="dept", parent=root, defaults={"field_type": ModelLabelField.FieldChoices.DATA, "label": "部门"}
        )
        dept = DeptInfo.objects.create(name="研发部", code="dev")
        UserInfo.objects.create_user(username="fkuser", password="Test@123456", dept=dept)
        ModelLabelField.objects.get_or_create(
            name="date_joined",
            parent=root,
            defaults={"field_type": ModelLabelField.FieldChoices.DATA, "label": "加入时间"},
        )
        dataset.columns = ["username", "dept", "date_joined"]
        dataset.save()

        report = _make_report(dataset, superuser, recipients=["boss@corp.com"])
        response = auth_client.post(f"{REPORT_URL}/{report.pk}/run", {}, format="json")
        assert response.status_code == 200, response.data
        task_id = response.json()["data"]["task_id"]

        record = ExportRecord.objects.get(pk=task_id)
        assert record.status == ExportRecord.Status.SUCCESS
        sheet = load_workbook(record.file.filepath).active
        assert sheet.cell(row=2, column=1).value == "fkuser"
        assert sheet.cell(row=2, column=2).value == str(dept.pk)
        joined = sheet.cell(row=2, column=3).value
        assert isinstance(joined, datetime) and joined.tzinfo is None

    def test_email_failure_degrades(self, auth_client, dataset, superuser, settings):
        """邮件失败：产物 SUCCESS 保留，last_status 标记邮件错误。"""
        settings.EMAIL_BACKEND = "django.core.mail.backends.locmem.EmailBackend"
        report = _make_report(dataset, superuser, recipients=["boss@corp.com"])
        from unittest import mock

        with mock.patch("system.analysis_tasks.EmailMessage.send", side_effect=Exception("smtp down")):
            auth_client.post(f"{REPORT_URL}/{report.pk}/run", {}, format="json")
        report.refresh_from_db()
        assert report.last_status == "SUCCESS_WITH_EMAIL_ERROR"


# ---------------------------------------------------------------- helpers


def _make_report(dataset, creator, frequency="daily", send_time="08:00", weekday=0, recipients=None, mode="rows"):
    return Report.objects.create(
        name=f"报表-{dataset.name}-{frequency}-{creator.username}",
        dataset=dataset,
        mode=mode,
        frequency=frequency,
        send_time=send_time,
        weekday=weekday,
        recipients=recipients if recipients is not None else ["a@corp.com"],
        creator=creator,
    )


def _second_user(screen_urls):
    from rest_framework.test import APIClient

    user = UserInfo.objects.create_user(username="lisi", password="Test@123456")
    grant(user, screen_urls)
    client = APIClient(HTTP_USER_AGENT="pytest-agent")
    client.force_authenticate(user=user)
    return {"user": user, "client": client}


@pytest.fixture
def dashboard_factory(superuser):
    def _make(name="大屏用看板"):
        return Dashboard.objects.create(name=name, creator=superuser, visibility="shared")

    return _make
