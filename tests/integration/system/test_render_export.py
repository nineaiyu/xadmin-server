# -*- coding: utf-8 -*-
"""导出渲染器 API 级测试：export-data 端点触发 xlsx / csv / import / update 模板。

渲染链路：export-data action 依据 ?type= 设置 format_kwarg 选择渲染器，
并经 BaseViewSet.paginate_queryset 关闭分页；本文件从接口层验证
Content-Type、Content-Disposition、表头、数据行与特殊字符处理。
"""

import codecs
import io

import openpyxl
import pytest

pytestmark = pytest.mark.django_db

ROLE_URL = "/api/system/role"
EXPORT_URL = f"{ROLE_URL}/export-data"
MENU_EXPORT_URL = "/api/system/menu/export-data"


def _create_role(auth_client, name="导出角色", code="export_role", **kwargs):
    payload = {"name": name, "code": code, "fields": {}}
    payload.update(kwargs)
    resp = auth_client.post(ROLE_URL, payload, format="json")
    assert resp.data["code"] == 1000, resp.data
    return resp.data["data"]["pk"]


def _load_xlsx(resp):
    return openpyxl.load_workbook(io.BytesIO(resp.content)).active


class TestExportXlsx:
    def test_xlsx_export_basic(self, auth_client, menu_factory):
        """带关联数据的 xlsx 导出：表头/数据行/布尔/关联字段渲染。"""
        menu = menu_factory("导出菜单", path="api/demo/book$", method="GET")
        _create_role(auth_client, name="导出角色A", code="role_x1", menu=[menu.pk])
        resp = auth_client.get(f"{EXPORT_URL}?type=xlsx")

        assert resp.status_code == 200
        assert resp["Content-Type"].startswith("application/xlsx")
        disposition = resp["Content-Disposition"]
        # 文件名前缀取模型名小写（UserRole → userrole）
        assert 'filename="userrole_' in disposition and '.xlsx"' in disposition
        assert resp["Access-Control-Expose-Headers"] == "Content-Disposition"

        ws = _load_xlsx(resp)
        headers = [c.value for c in ws[1]]
        # 首列为主键，标题形如 xxx(pk)；业务列表头形如 xxx(name)
        assert any(str(h).endswith("(pk)") for h in headers)
        assert any(str(h).endswith("(name)") for h in headers)

        rows = list(ws.iter_rows(min_row=2, values_only=True))
        assert rows, "导出应包含数据行"
        joined = ",".join(str(v) for v in rows[0])
        assert "导出角色A" in joined
        assert "role_x1" in joined
        # 布尔字段渲染为 Yes/No
        assert "Yes" in joined or "No" in joined
        # 关联菜单渲染为 name(pk) 形式
        assert "导出菜单(" in joined and str(menu.pk) in joined

    def test_xlsx_export_empty_dataset(self, auth_client):
        """空数据集导出：仅表头行，仍是合法 xlsx。"""
        resp = auth_client.get(f"{EXPORT_URL}?type=xlsx&name=不存在的角色")
        assert resp.status_code == 200
        assert resp["Content-Type"].startswith("application/xlsx")
        assert _load_xlsx(resp).max_row == 1

    def test_export_with_chinese_and_special_chars(self, auth_client):
        """中文与特殊字符（换行/等号）不破坏 xlsx 结构。"""
        _create_role(auth_client, name="中文=引号\n换行", code="role_cn1")
        resp = auth_client.get(f"{EXPORT_URL}?type=xlsx")
        ws = _load_xlsx(resp)
        values = [str(v) for row in ws.iter_rows(min_row=2, values_only=True) for v in row]
        assert any("中文" in v and "换行" in v for v in values)


class TestExportCsv:
    def test_csv_export_basic_and_escape(self, auth_client):
        """csv 导出：BOM 头、表头、以转义字符开头的数据被前置单引号。"""
        _create_role(auth_client, name="=evil角色", code="role_c1")
        _create_role(auth_client, name="普通角色", code="role_c2")
        resp = auth_client.get(f"{EXPORT_URL}?type=csv")

        assert resp.status_code == 200
        assert resp["Content-Type"].startswith("text/csv")
        assert resp.content.startswith(codecs.BOM_UTF8)
        text = resp.content.decode("utf-8-sig")
        assert "(name)" in text.splitlines()[0]
        # CSV_FILE_ESCAPE_CHARS（=/@/0）开头的单元格需转义，防公式注入
        assert "'=evil角色" in text
        assert "普通角色" in text

    def test_csv_export_empty_dataset(self, auth_client):
        resp = auth_client.get(f"{EXPORT_URL}?type=csv&name=不存在")
        assert resp.status_code == 200
        assert resp["Content-Type"].startswith("text/csv")
        text = resp.content.decode("utf-8-sig")
        assert len(text.splitlines()) == 1


class TestExportTemplates:
    def test_import_template_has_help_row_and_template_filename(self, auth_client):
        """import 模板：文件名后缀为 template，第二行为 #Help 说明行。"""
        _create_role(auth_client)
        resp = auth_client.get(f"{EXPORT_URL}?type=xlsx&template=import")
        assert resp.status_code == 200
        assert 'userrole_template.xlsx"' in resp["Content-Disposition"]
        ws = _load_xlsx(resp)
        assert str(ws.cell(row=2, column=1).value).startswith("#Help")

    def test_update_template_includes_pk_and_help_row(self, auth_client):
        """update 模板：首列带主键，且携带 #Help 说明行。"""
        _create_role(auth_client)
        resp = auth_client.get(f"{EXPORT_URL}?type=xlsx&template=update")
        assert resp.status_code == 200
        ws = _load_xlsx(resp)
        headers = [c.value for c in ws[1]]
        assert any(str(h).endswith("(pk)") for h in headers)
        assert str(ws.cell(row=2, column=1).value).startswith("#Help")

    def test_unknown_type_falls_back_to_first_renderer(self, auth_client):
        """未知 type：内容协商强制回落到第一个渲染器（xlsx）。"""
        _create_role(auth_client)
        resp = auth_client.get(f"{EXPORT_URL}?type=unknown")
        assert resp.status_code == 200
        assert resp["Content-Type"].startswith("application/xlsx")

    def test_menu_import_template_includes_pk_for_self_relation(self, auth_client):
        """菜单（自关联模型）import 模板：需包含 pk 列用于拓扑排序。"""
        resp = auth_client.get(f"{MENU_EXPORT_URL}?type=xlsx&template=import")
        assert resp.status_code == 200, resp.data
        headers = [str(c.value) for c in _load_xlsx(resp)[1]]
        assert any(h.endswith("(pk)") for h in headers)
