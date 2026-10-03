# -*- coding: utf-8 -*-
"""报表设计的投递渲染集成测试（批次二）。

覆盖：空 design 保持存量单表口径；design.columns 裁剪明细列且保序；table_limit 截断明细行；
聚合组件各落一张独立 sheet；单组件失败 fail-soft（写提示行，不拖垮整份报表）。
"""

import io

import pytest
from django.utils.translation import gettext as _t
from openpyxl import load_workbook

from dataset.analysis_tasks import _render_workbook
from dataset.models.dataset import Dataset, Report
from system.models import ModelLabelField

pytestmark = pytest.mark.django_db

ALL_COLUMNS = ["username", "gender", "is_active"]


@pytest.fixture
def model_registry(db):
    """数据集可用字段白名单：执行侧只认已登记的数据字段（与 test_analysis_api 同款）。"""
    root, _ = ModelLabelField.objects.get_or_create(
        name="system.userinfo", defaults={"field_type": ModelLabelField.FieldChoices.DATA, "label": "用户"}
    )
    for name in ALL_COLUMNS:
        ModelLabelField.objects.get_or_create(
            name=name, parent=root, defaults={"field_type": ModelLabelField.FieldChoices.DATA, "label": name}
        )
    return root


@pytest.fixture
def dataset(model_registry, superuser):
    return Dataset.objects.create(
        name="投递设计数据集",
        bound_model="system.userinfo",
        columns=list(ALL_COLUMNS),
        visibility="shared",
        creator=superuser,
    )


def _workbook(dataset, superuser, design, mode="rows", **kwargs):
    report = Report.objects.create(
        name=f"投递报表-{design.get('table_limit', 'x')}-{mode}-{len(design.get('components', []))}",
        dataset=dataset,
        recipients=["deliver@example.com"],
        design=design,
        mode=mode,
        **kwargs,
    )
    from system.models import UserInfo

    user = UserInfo.objects.get(pk=superuser.pk)
    content, rows = _render_workbook(report, user)
    return load_workbook(io.BytesIO(content)), rows


def test_legacy_design_keeps_single_sheet(dataset, superuser):
    wb, rows = _workbook(dataset, superuser, {})
    assert wb.sheetnames == [_t("Report")]
    sheet = wb[_t("Report")]
    assert [cell.value for cell in sheet[1]] == ALL_COLUMNS
    assert rows == sheet.max_row - 1


def test_design_columns_trim_and_order(dataset, superuser):
    wb, rows = _workbook(dataset, superuser, {"columns": ["gender", "username"], "table_limit": 10})
    sheet = wb[_t("Report")]
    assert [cell.value for cell in sheet[1]] == ["gender", "username"]
    assert rows <= 10


def test_table_limit_truncates_rows(dataset, superuser):
    from system.models import UserInfo

    for index in range(12):
        UserInfo.objects.create_user(username=f"design_row_{index}", password="Pass-2026!")
    wb, rows = _workbook(dataset, superuser, {"columns": ["username"], "table_limit": 10})
    assert rows <= 10


def test_component_sheets(dataset, superuser):
    wb, _ = _workbook(
        dataset,
        superuser,
        {
            "columns": ["username"],
            "components": [
                {"id": "c1", "type": "bar", "title": "性别分布", "group_by": "gender", "metric": "count"},
                {"id": "n1", "type": "number", "title": "性别合计", "metric": "sum", "value_field": "gender"},
            ],
        },
    )
    assert wb.sheetnames[0] == _t("Report")
    assert len(wb.sheetnames) == 3
    # 组件 sheet 名带序号前缀（唯一 + 截断），表头为「名称 / 值」
    titles = wb.sheetnames[1:]
    assert titles[0].startswith("1-性别分布") and titles[1].startswith("2-性别合计")
    sheet = wb[titles[0]]
    assert [cell.value for cell in sheet[1]] == [_t("Name"), _t("Value")]
    assert sheet.max_row >= 2


def test_aggregate_mode_with_components_keeps_legacy_sheet(dataset, superuser):
    wb, rows = _workbook(
        dataset,
        superuser,
        {"components": [{"id": "c1", "type": "pie", "group_by": "is_active", "metric": "count"}]},
        mode="aggregate",
        group_by="gender",
        metric="count",
    )
    assert wb.sheetnames[0] == _t("Report")
    sheet = wb[_t("Report")]
    assert [cell.value for cell in sheet[1]] == [_t("Name"), _t("Value")]
    assert rows == sheet.max_row - 1
    assert len(wb.sheetnames) == 2


def test_component_failure_is_fail_soft(dataset, superuser, monkeypatch):
    import dataset.utils.dataset as dataset_utils

    def boom(*args, **kwargs):
        raise ValueError("field gone")

    monkeypatch.setattr(dataset_utils, "aggregate_dataset", boom)
    wb, _ = _workbook(
        dataset,
        superuser,
        {"components": [{"id": "c1", "type": "bar", "title": "坏组件", "group_by": "gender"}]},
    )
    assert len(wb.sheetnames) == 2
    sheet = wb[wb.sheetnames[1]]
    assert [cell.value for cell in sheet[1]] == [_t("Name"), _t("Value")]
    assert sheet.max_row == 2
    assert sheet.cell(row=2, column=1).value == _t("Chart data unavailable")
