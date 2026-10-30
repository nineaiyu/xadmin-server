# -*- coding: utf-8 -*-
"""数据集执行链的值级脱敏：REST / 定时报表 / 仪表盘与大屏共用同一内核。"""

from io import BytesIO

import pytest
from django.core.exceptions import ValidationError
from openpyxl import load_workbook

from audit.models import DataMaskRule
from dataset.dashboard_data import build_dashboard_data_payload
from dataset.models.dataset import Dashboard, Dataset, Report, Screen
from dataset.screen_data import build_screen_data_payload
from dataset.utils.dataset import aggregate_dataset, execute_dataset
from identity.models import UserInfo, UserRole
from system.models import DataPermission, ModelLabelField

pytestmark = pytest.mark.django_db

MASKED_NAME = "张*丰"


@pytest.fixture(autouse=True)
def _data_permission_on(settings):
    settings.PERMISSION_DATA_ENABLED = True


@pytest.fixture
def model_registry(db):
    root, _ = ModelLabelField.objects.get_or_create(
        name="identity.userinfo",
        defaults={"field_type": ModelLabelField.FieldChoices.DATA, "label": "用户"},
    )
    for name in ("username", "nickname", "email", "id"):
        ModelLabelField.objects.get_or_create(
            name=name,
            parent=root,
            defaults={"field_type": ModelLabelField.FieldChoices.DATA, "label": name},
        )
    return root


@pytest.fixture
def dataset(model_registry, superuser):
    return Dataset.objects.create(
        name="掩码数据集",
        bound_model="identity.userinfo",
        columns=["username", "nickname", "email"],
        visibility="shared",
        creator=superuser,
    )


@pytest.fixture
def masked_user(db):
    user = UserInfo.objects.create_user(
        username="masked-viewer", password="Test@123456", nickname="张三丰", email="masked@example.com"
    )
    role = UserRole.objects.create(name="mask-role", code="mask-role")
    user.roles.add(role)
    rule = DataMaskRule.objects.create(
        model="identity.userinfo",
        field="nickname",
        mask_type=DataMaskRule.MaskType.NAME,
        keep_head=1,
        keep_tail=1,
    )
    rule.roles.add(role)
    # 行级数据权限：仅本人行（否则 fail-closed 空集，掩码断言无从谈起）
    row_rule = DataPermission.objects.create(
        name="dp-masked-viewer-self",
        rules=[
            {
                "table": "identity.userinfo",
                "field": "id",
                "type": "value.user.id",
                "match": "exact",
                "value": "",
                "exclude": False,
            }
        ],
    )
    user.rules.add(row_rule)
    return user


class TestExecuteDatasetMask:
    def test_non_superuser_rows_are_masked(self, dataset, masked_user):
        result = execute_dataset(dataset, masked_user)
        assert result["rows"]
        assert result["rows"][0]["nickname"] == MASKED_NAME

    def test_superuser_keeps_original(self, dataset, superuser, masked_user):
        result = execute_dataset(dataset, superuser)
        assert result["rows"]
        assert any(row["nickname"] == "张三丰" for row in result["rows"])


class TestAggregateDatasetMask:
    def test_group_values_are_masked(self, dataset, masked_user):
        result = aggregate_dataset(dataset, masked_user, group_by="nickname", metric="count")
        assert result["series"]
        assert all(item["name"] == MASKED_NAME for item in result["series"])
        assert all(item["raw"] == MASKED_NAME for item in result["series"])

    def test_sum_on_masked_field_fails_closed(self, dataset, masked_user, model_registry):
        """聚合值无法掩码：已掩码字段不得被 sum/avg（否则原值分布直接泄露）。"""
        rule = DataMaskRule.objects.create(
            model="identity.userinfo",
            field="id",
            mask_type=DataMaskRule.MaskType.CUSTOM,
            pattern=".",
        )
        rule.roles.add(*masked_user.roles.all())
        dataset.columns = ["username", "id"]
        dataset.save(update_fields=["columns"])
        with pytest.raises(ValidationError):
            aggregate_dataset(dataset, masked_user, group_by="", metric="sum", value_field="id")


class TestWsPushMask:
    def test_dashboard_payload_masks_group_values(self, dataset, masked_user):
        dashboard = Dashboard.objects.create(
            name="掩码看板",
            visibility="shared",
            creator=masked_user,
            layout=[
                {
                    "id": "card-1",
                    "dataset": str(dataset.pk),
                    "chart_type": "bar",
                    "group_by": "nickname",
                    "metric": "count",
                }
            ],
        )
        payload = build_dashboard_data_payload(masked_user, dashboard)
        data = payload["cards"][0]["data"]
        assert all(item["name"] == MASKED_NAME for item in data["series"])
        assert all(item["raw"] == MASKED_NAME for item in data["series"])

    def test_screen_payload_masks_group_values(self, dataset, masked_user):
        dashboard = Dashboard.objects.create(
            name="掩码大屏看板",
            visibility="shared",
            creator=masked_user,
            layout=[
                {
                    "id": "card-1",
                    "dataset": str(dataset.pk),
                    "chart_type": "bar",
                    "group_by": "nickname",
                    "metric": "count",
                }
            ],
        )
        screen = Screen.objects.create(
            name="掩码大屏",
            dashboards=[str(dashboard.pk)],
            visibility="shared",
            creator=masked_user,
        )
        frames = build_screen_data_payload(masked_user, screen, rev=1)
        data = frames[0]["cards"][0]["data"]
        assert all(item["name"] == MASKED_NAME for item in data["series"])
        assert all(item["raw"] == MASKED_NAME for item in data["series"])


class TestScheduledReportMask:
    def test_workbook_cells_are_masked(self, dataset, masked_user):
        from dataset.report_render import _render_workbook

        report = Report.objects.create(
            name="掩码报表",
            dataset=dataset,
            recipients=["ops@example.com"],
            creator=masked_user,
        )
        content, rows = _render_workbook(report, masked_user)
        assert rows >= 1
        sheet = load_workbook(BytesIO(content)).active
        headers = [cell.value for cell in sheet[1]]
        nickname_index = headers.index("nickname") + 1
        values = [sheet.cell(row=index, column=nickname_index).value for index in range(2, sheet.max_row + 1)]
        assert MASKED_NAME in values
        assert "张三丰" not in values

    def test_aggregate_workbook_cells_are_masked(self, dataset, masked_user):
        from dataset.report_render import _render_workbook

        report = Report.objects.create(
            name="聚合掩码报表",
            dataset=dataset,
            mode="aggregate",
            group_by="nickname",
            metric="count",
            recipients=["ops@example.com"],
            creator=masked_user,
        )
        content, _rows = _render_workbook(report, masked_user)
        sheet = load_workbook(BytesIO(content)).active
        values = [sheet.cell(row=index, column=1).value for index in range(2, sheet.max_row + 1)]
        assert MASKED_NAME in values
        assert "张三丰" not in values
