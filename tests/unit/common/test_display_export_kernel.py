# -*- coding: utf-8 -*-
"""内核显示名与导出装配 API：user_label / excel_safe_value / to_xlsx / save_xlsx。"""

import datetime
import uuid
from io import BytesIO
from zoneinfo import ZoneInfo

from django.utils import timezone
from django.utils.translation import gettext_lazy as _

from common.utils.display import user_label
from common.utils.excel import excel_safe_value, save_xlsx, to_xlsx


class _User:
    def __init__(self, nickname=None, username=None):
        self.nickname = nickname
        self.username = username


class TestUserLabel:
    def test_nickname_preferred_over_username(self):
        assert user_label(_User(nickname=" 昵称 ", username="u1")) == "昵称"

    def test_falls_back_to_username(self):
        assert user_label(_User(nickname="", username="u1")) == "u1"
        assert user_label(_User(nickname=None, username="u1")) == "u1"

    def test_none_and_empty_are_safe(self):
        assert user_label(None) == ""
        assert user_label(_User(), default="-") == "-"
        # 鸭子对象：缺属性不抛错
        assert user_label(object(), default="x") == "x"


class TestExcelSafeValue:
    def test_uuid_and_lazy_to_str(self):
        assert excel_safe_value(uuid.UUID(int=42)) == str(uuid.UUID(int=42))
        assert excel_safe_value(_("lazy")) == "lazy"

    def test_aware_datetime_and_time_to_local_naive(self):
        aware = timezone.make_aware(datetime.datetime(2026, 1, 1, 12, 0))
        naive = excel_safe_value(aware)
        assert isinstance(naive, datetime.datetime)
        assert naive.tzinfo is None
        assert naive == timezone.localtime(aware).replace(tzinfo=None)
        aware_time = datetime.time(8, 0, tzinfo=ZoneInfo("UTC"))
        assert excel_safe_value(aware_time).tzinfo is None

    def test_containers_to_json_text(self):
        assert excel_safe_value({"a": 1}) == '{"a": 1}'
        assert excel_safe_value([1, "x"]) == '[1, "x"]'

    def test_scalars_pass_through(self):
        assert excel_safe_value(None) is None
        assert excel_safe_value(3.14) == 3.14
        assert excel_safe_value(True) is True


class TestToXlsx:
    SHEETS = {"汇总": [["名称", "数值"], ["=HYPERLINK(...)", 10]], "明细": [["列"], ["值"]]}

    def test_multi_sheet_bytes_readable_by_openpyxl(self):
        from openpyxl import load_workbook

        content = to_xlsx(self.SHEETS)
        wb = load_workbook(BytesIO(content))
        assert wb.sheetnames == ["汇总", "明细"]
        assert wb["汇总"]["A2"].data_type == "s"  # 公式样文本被强制为文本（防注入）
        assert wb["汇总"]["B2"].value == 10  # 数值保持类型
        assert wb["汇总"]["A1"].font.bold is True  # 表头加粗

    def test_lazy_header_and_options(self):
        from openpyxl import load_workbook

        content = to_xlsx({_("报表"): [[_("名称")], ["a"]]}, bold_header=False, auto_width=False)
        wb = load_workbook(BytesIO(content))
        assert wb.active.title == "报表"  # 惰性翻译表头可装配
        assert wb.active["A1"].value == "名称"
        assert not wb.active["A1"].font.bold  # 关闭加粗（openpyxl 缺省字体 b=False）

    def test_save_xlsx_writes_file(self, tmp_path):
        from openpyxl import load_workbook

        target = tmp_path / "report.xlsx"
        save_xlsx(target, self.SHEETS)
        assert target.exists()
        assert load_workbook(target).sheetnames == ["汇总", "明细"]
