# -*- coding: utf-8 -*-
"""监控增强能力：历史趋势聚合/环比、健康总览、告警记录、事件查询、阈值设置与报表导出。"""

import datetime
import io
import uuid

import pytest
from django.test import override_settings
from django.utils import timezone
from rest_framework.test import APIRequestFactory, force_authenticate

from common.models import Monitor, MonitorAlert
from system.models.log import OperationLog
from system.models.task import TaskExecution
from system.views.monitor import MonitorViewSet

pytestmark = pytest.mark.django_db

THRESHOLD_SETTINGS = {
    "SECURITY_MONITOR_CPU_PERCENT_MAX": 80,
    "SECURITY_MONITOR_CPU_LOAD_MAX": 5,
    "SECURITY_MONITOR_MEMORY_USED_MAX": 85,
    "SECURITY_MONITOR_DISK_USED_MAX": 80,
}
DEFAULTS = {"cpu_load": 0.5, "cpu_percent": 10.0, "memory_used": 20.0, "disk_used": 30.0}
ACTION_MAP = {
    "overview": "overview",
    "history": "history",
    "thresholds": "thresholds",
    "events": "events",
    "export": "export",
    "services": "services",
}

ACTIONMETHOD_MAP = {
    ("thresholds", "get"): "thresholds",
    ("thresholds", "put"): "thresholds",
    ("history", "get"): "history",
    ("events", "get"): "events",
    ("export", "get"): "export",
}


def invoke(action, method="get", user=None, query="", data=None):
    factory = APIRequestFactory()
    path = f"/api/system/monitor/{action}"
    if query:
        path = f"{path}?{query}"
    if method == "get":
        request = factory.get(path)
    else:
        request = getattr(factory, method)(path, data or {}, format="json")
    if user:
        force_authenticate(request, user=user)
    view_action = ACTIONMETHOD_MAP[(action, method)]
    return MonitorViewSet.as_view({method: view_action})(request)


def seed_monitor(points):
    """points: [(datetime, fields)]；created_time 为 auto_now_add，需 update 覆盖。"""
    for timestamp, values in points:
        row = Monitor.objects.create(**{**DEFAULTS, **values})
        Monitor.objects.filter(pk=row.pk).update(created_time=timestamp)


def _slot(offset_minutes, extra_seconds=0):
    """对齐到 5 分钟桶起点的时间（避免跨桶边界导致断言脆弱）。"""
    base = int((timezone.now() - datetime.timedelta(minutes=offset_minutes)).timestamp())
    slot = base // 300 * 300
    return datetime.datetime.fromtimestamp(slot + extra_seconds, tz=datetime.UTC)


class TestHistory:
    def test_buckets_metrics_and_previous_compare(self, superuser):
        # 当前窗口：同一 5m 桶两条 + 下一桶一条；上一窗口（>1h 前）一条供环比
        seed_monitor(
            [
                (_slot(20, 60), {"cpu_percent": 10.0}),
                (_slot(20, 120), {"cpu_percent": 30.0}),
                (_slot(15, 60), {"cpu_percent": 50.0}),
                (_slot(70), {"cpu_percent": 40.0}),
            ]
        )
        resp = invoke("history", user=superuser, query="range=1h&interval=5m&metrics=cpu_percent,memory_used")
        assert resp.data["code"] == 1000
        data = resp.data["data"]
        assert data["range"]["interval"] == "5m"
        assert data["metrics"] == ["cpu_percent", "memory_used"]
        # 同桶两条取均值，不同桶分开
        cpu_values = [point["cpu_percent"] for point in data["points"]]
        assert 20.0 in cpu_values and 50.0 in cpu_values
        assert len(data["points"]) == 2
        # 汇总：窗口内全部原始点
        assert data["summary"]["cpu_percent"]["avg"] == pytest.approx(30.0)
        assert data["summary"]["cpu_percent"]["max"] == 50.0
        assert data["summary"]["cpu_percent"]["min"] == 10.0
        # 环比：上一等长窗口的均值
        assert data["compare"]["cpu_percent"]["prev_avg"] == pytest.approx(40.0)
        assert data["compare"]["cpu_percent"]["delta"] == pytest.approx(-10.0)

    def test_net_rate_derived_from_counters(self, superuser):
        now = timezone.now()
        seed_monitor(
            [
                (now - datetime.timedelta(minutes=10), {"net_sent_mb": 100.0, "net_recv_mb": 200.0}),
                (now - datetime.timedelta(minutes=9), {"net_sent_mb": 160.0, "net_recv_mb": 260.0}),
            ]
        )
        resp = invoke("history", user=superuser, query="range=1h&interval=1m&metrics=net_sent_rate,net_recv_rate")
        rates = [point.get("net_sent_rate") for point in resp.data["data"]["points"]]
        # 60MB / 60s = 1MB/s = 1024KB/s
        assert 1024.0 in rates

    def test_counter_reset_marks_rate_missing(self, superuser):
        now = timezone.now()
        seed_monitor(
            [
                (now - datetime.timedelta(minutes=10), {"net_sent_mb": 500.0}),
                (now - datetime.timedelta(minutes=9), {"net_sent_mb": 10.0}),
            ]
        )
        resp = invoke("history", user=superuser, query="range=1h&interval=1m&metrics=net_sent_rate")
        data = resp.data["data"]
        # 负差值（计数器重置）按缺失处理：不出现在聚合点里，汇总 last 为 None
        assert all("net_sent_rate" not in point for point in data["points"])
        assert data["summary"]["net_sent_rate"]["last"] is None

    def test_unknown_metric_falls_back_to_defaults(self, superuser):
        resp = invoke("history", user=superuser, query="range=1h&metrics=not_exist")
        assert resp.data["data"]["metrics"] == ["cpu_percent", "memory_used", "disk_used"]

    def test_auto_interval_by_window(self, superuser):
        resp = invoke("history", user=superuser, query="range=24h")
        assert resp.data["data"]["range"]["interval"] == "15m"
        resp = invoke("history", user=superuser, query="range=30d")
        assert resp.data["data"]["range"]["interval"] == "1d"

    def test_explicit_start_end_window(self, superuser):
        resp = invoke("history", user=superuser, query="start=2026-09-18T00:00:00&end=2026-09-18T06:00:00")
        data = resp.data["data"]
        assert data["range"]["range_key"] == "custom"
        assert data["range"]["start"].startswith("2026-09-18T00:00")


class TestHealthSummary:
    def test_critical_resource_and_degraded_celery(self):
        from system.utils import metrics

        health = metrics.collect_health_summary(
            live={"cpu_percent": 100, "cpu_load": 0.1, "memory_used": 10, "disk_used": 10},
            services={
                "db": {"status": True, "cost": 0.01},
                "redis": {"status": True, "cost": 0.01},
                "celery": {"status": False, "cost": 0.0},
            },
            celery_skipped=False,
        )
        by_key = {item["key"]: item for item in health["items"]}
        # 100 >= 80 * 1.2 → critical
        assert by_key["cpu_percent"]["status"] == "critical"
        # 未探测到 worker 降为 warning（不阻断核心服务）
        assert by_key["celery"]["status"] == "warning"
        assert health["status"] == "critical"
        assert health["score"] < 100

    def test_all_healthy_scores_full(self):
        from system.utils import metrics

        health = metrics.collect_health_summary(
            live={"cpu_percent": 10, "cpu_load": 0.5, "memory_used": 20, "disk_used": 30},
            services={
                "db": {"status": True, "cost": 0.01},
                "redis": {"status": True, "cost": 0.01},
                "celery": {"status": True, "cost": 0.01},
            },
            celery_skipped=False,
        )
        assert health["status"] == "healthy"
        assert health["score"] == 100

    def test_skipped_celery_not_counted(self):
        from system.utils import metrics

        health = metrics.collect_health_summary(
            live={"cpu_percent": 10, "cpu_load": 0.5, "memory_used": 20, "disk_used": 30},
            services={"db": {"status": True}, "redis": {"status": True}, "celery": {"status": False}},
            celery_skipped=True,
        )
        by_key = {item["key"]: item for item in health["items"]}
        assert by_key["celery"]["status"] == "unknown"
        assert health["status"] == "healthy"


class TestThresholds:
    def test_read_returns_serializer_metadata(self, superuser):
        with override_settings(**THRESHOLD_SETTINGS):
            resp = invoke("thresholds", user=superuser)
        data = resp.data["data"]
        items = {item["key"]: item for item in data["items"]}
        assert set(items) == set(THRESHOLD_SETTINGS)
        assert items["SECURITY_MONITOR_DISK_USED_MAX"]["value"] == 80
        assert items["SECURITY_MONITOR_DISK_USED_MAX"]["max"] == 100
        assert data["check_interval_seconds"] == 60

    def test_update_persists_and_takes_effect(self, superuser):
        from settings.models import Setting

        with override_settings(**THRESHOLD_SETTINGS):
            resp = invoke("thresholds", method="put", user=superuser, data={"SECURITY_MONITOR_DISK_USED_MAX": 90})
            assert resp.data["code"] == 1000
            assert resp.data["data"]["changed"] == ["SECURITY_MONITOR_DISK_USED_MAX"]
            setting = Setting.objects.get(name="SECURITY_MONITOR_DISK_USED_MAX")
            assert setting.cleaned_value == 90
            # 本进程立即生效（refresh_setting），无需等待 pubsub
            follow = invoke("thresholds", user=superuser)
            items = {item["key"]: item for item in follow.data["data"]["items"]}
            assert items["SECURITY_MONITOR_DISK_USED_MAX"]["value"] == 90

    def test_update_rejects_out_of_range(self, superuser):
        with override_settings(**THRESHOLD_SETTINGS):
            resp = invoke("thresholds", method="put", user=superuser, data={"SECURITY_MONITOR_DISK_USED_MAX": 101})
        assert resp.status_code == 400


class TestAlerts:
    def test_state_transitions_firing_then_resolved(self, superuser):
        from common.notifications import ServerPerformanceCheckUtil

        with override_settings(**THRESHOLD_SETTINGS):
            seed_monitor([(timezone.now(), {"disk_used": 95})])
            util = ServerPerformanceCheckUtil()
            util.check()
            util.sync_alert_records()
            alert = MonitorAlert.objects.get()
            assert alert.item == "disk_used"
            assert alert.status == MonitorAlert.Status.FIRING
            assert alert.count == 1

            # 持续超标：续写不新增
            util = ServerPerformanceCheckUtil()
            util.check()
            util.sync_alert_records()
            assert MonitorAlert.objects.count() == 1
            alert.refresh_from_db()
            assert alert.count == 2

            # 指标回落：置为已恢复
            Monitor.objects.all().delete()
            seed_monitor([(timezone.now(), {"disk_used": 30})])
            util = ServerPerformanceCheckUtil()
            util.check()
            util.sync_alert_records()
            alert.refresh_from_db()
            assert alert.status == MonitorAlert.Status.RESOLVED
            assert alert.resolved_time is not None


class TestEvents:
    def test_alert_error_task_kinds(self, superuser):
        now = timezone.now()
        MonitorAlert.objects.create(
            item="disk_used", value=95, threshold=80, message="disk", first_time=now, last_time=now
        )
        OperationLog.objects.create(module="demo", status_code=1001, exec_time=0.2, creator=superuser)
        TaskExecution.objects.create(
            id=uuid.uuid4(),
            name="demo.task",
            status=TaskExecution.Status.FAILURE,
            date_start=now,
            date_finished=now,
            creator=superuser,
        )

        alerts = invoke("events", user=superuser, query="kind=alert").data["data"]
        assert len(alerts["results"]) == 1
        assert alerts["counts"]["firing"] == 1

        errors = invoke("events", user=superuser, query="kind=error&range=24h").data["data"]
        assert [row["status_code"] for row in errors["results"]] == [1001]

        tasks = invoke("events", user=superuser, query="kind=task&range=24h").data["data"]
        assert [row["name"] for row in tasks["results"]] == ["demo.task"]

    def test_alert_filter_by_status_and_item(self, superuser):
        now = timezone.now()
        MonitorAlert.objects.create(item="disk_used", value=95, threshold=80, first_time=now, last_time=now)
        MonitorAlert.objects.create(
            item="cpu_percent",
            status=MonitorAlert.Status.RESOLVED,
            value=99,
            threshold=80,
            first_time=now,
            last_time=now,
            resolved_time=now,
        )
        firing = invoke("events", user=superuser, query="kind=alert&status=firing").data["data"]
        assert [row["item"] for row in firing["results"]] == ["disk_used"]
        by_item = invoke("events", user=superuser, query="kind=alert&item=cpu_percent").data["data"]
        assert [row["item"] for row in by_item["results"]] == ["cpu_percent"]


class TestExport:
    def test_history_csv_and_xlsx(self, superuser):
        from system.utils.monitor_history import metric_label

        seed_monitor([(timezone.now(), {"cpu_percent": 42.0})])
        resp = invoke("export", user=superuser, query="kind=history&type=csv&range=1h&metrics=cpu_percent")
        assert resp.status_code == 200
        assert "attachment" in resp["Content-Disposition"]
        content = resp.content.decode("utf-8-sig")
        assert metric_label("cpu_percent") in content
        assert "42" in content

        resp = invoke("export", user=superuser, query="kind=history&type=xlsx&range=1h&metrics=cpu_percent")
        assert resp.content[:2] == b"PK"
        from openpyxl import load_workbook

        workbook = load_workbook(io.BytesIO(resp.content))
        assert len(workbook.sheetnames) == 2
        assert workbook[workbook.sheetnames[0]].max_row >= 2

    def test_alerts_csv_uses_same_origin_labels(self, superuser):
        now = timezone.now()
        MonitorAlert.objects.create(item="disk_used", value=95, threshold=80, first_time=now, last_time=now)
        resp = invoke("export", user=superuser, query="kind=alerts&type=csv")
        content = resp.content.decode("utf-8-sig")
        # 状态文案取模型 choices（本机有 .mo 显中文、CI 无 .mo 显英文，须同源断言）
        assert str(MonitorAlert.Status.FIRING.label) in content
        assert str(MonitorAlert.Item.DISK_USED.label) in content
