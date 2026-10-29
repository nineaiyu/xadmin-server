# -*- coding: utf-8 -*-
"""监控告警导出：单查询游标取数，CSV 分块流式、XLSX write_only 装配。"""

import io
from datetime import timedelta

import pytest
from django.utils import timezone

from system.models import MonitorAlert
from system.services.monitor import events
from system.services.monitor.export import render_streaming_export

pytestmark = pytest.mark.django_db


def _create_alert(index, message=""):
    now = timezone.now()
    return MonitorAlert.objects.create(
        item=MonitorAlert.Item.CPU_PERCENT,
        status=MonitorAlert.Status.FIRING,
        value=90 + index,
        threshold=80,
        message=message or f"alert-{index}",
        count=index + 1,
        first_time=now - timedelta(minutes=5),
        last_time=now - timedelta(minutes=index),
    )


class TestAlertExport:
    def test_csv_streams_all_rows(self):
        for index in range(3):
            _create_alert(index)
        title, header = events.alert_export_columns()
        response = render_streaming_export("monitor-alerts", title, header, events.iter_alert_export_rows(), "csv")
        body = b"".join(response.streaming_content).decode("utf-8-sig")
        lines = [line for line in body.splitlines() if line]
        assert len(lines) == 4  # 表头 + 3 行
        assert "alert-0" in body and "alert-2" in body
        assert response["Content-Disposition"].endswith('.csv"')

    def test_xlsx_uses_write_only_and_escapes_formulas(self):
        _create_alert(0, message='=HYPERLINK("http://evil")')
        title, header = events.alert_export_columns()
        response = render_streaming_export("monitor-alerts", title, header, events.iter_alert_export_rows(), "xlsx")
        from openpyxl import load_workbook

        workbook = load_workbook(io.BytesIO(response.content))
        sheet = workbook.active
        assert sheet.max_row == 2
        message_cell = sheet.cell(row=2, column=len(header))
        assert message_cell.data_type == "s"
        assert message_cell.value.startswith("=")

    def test_iteration_is_single_query(self, django_assert_num_queries):
        for index in range(5):
            _create_alert(index)
        with django_assert_num_queries(1):
            rows = list(events.iter_alert_export_rows())
        assert len(rows) == 5
