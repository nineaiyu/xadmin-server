#!/usr/bin/env python
# -*- coding:utf-8 -*-
"""监控报表导出渲染（CSV / Excel 双格式，与下载中心同口径的通用表格落盘）。

调用方组装 sheets（标题 / 表头 / 行数据），本模块只负责编码与文件名。
CSV 带 UTF-8 BOM 供 Excel 正确识别中文；Excel 走 openpyxl，datetime 一律转
本地朴素时间（USE_TZ 下 aware datetime 会触发 openpyxl 拒绝写入）。
"""

import csv
import datetime
import io

from django.utils import timezone

CSV_CONTENT_TYPE = "text/csv; charset=utf-8"
XLSX_CONTENT_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def _excel_safe(value):
    """openpyxl 兼容转换：aware datetime → 本地朴素时间；其余原样。"""
    if isinstance(value, datetime.datetime) and timezone.is_aware(value):
        return timezone.localtime(value).replace(tzinfo=None)
    return value


def _csv_bytes(sheets):
    buffer = io.StringIO()
    writer = csv.writer(buffer)
    for index, sheet in enumerate(sheets):
        if index:
            writer.writerow([])
        if len(sheets) > 1:
            writer.writerow([sheet["title"]])
        writer.writerow(sheet["header"])
        for row in sheet["rows"]:
            writer.writerow(row)
    return buffer.getvalue().encode("utf-8-sig")


def _xlsx_bytes(sheets):
    from openpyxl import Workbook
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter

    workbook = Workbook()
    for index, sheet in enumerate(sheets):
        worksheet = workbook.active if index == 0 else workbook.create_sheet()
        worksheet.title = str(sheet["title"])[:31] or f"Sheet{index + 1}"
        worksheet.append(sheet["header"])
        for cell in worksheet[1]:
            cell.font = Font(bold=True)
        for row in sheet["rows"]:
            worksheet.append([_excel_safe(value) for value in row])
        for column, title in enumerate(sheet["header"], start=1):
            lengths = [len(str(title))]
            lengths += [len(str(row[column - 1])) for row in sheet["rows"] if column - 1 < len(row)]
            worksheet.column_dimensions[get_column_letter(column)].width = min(max(lengths) + 2, 40)
    stream = io.BytesIO()
    workbook.save(stream)
    return stream.getvalue()


def render_table_export(prefix, sheets, file_format="csv"):
    """生成导出文件，返回 (文件名, 内容 bytes, content-type)。"""
    timestamp = timezone.localtime(timezone.now()).strftime("%Y%m%d-%H%M%S")
    if file_format == "xlsx":
        return f"{prefix}-{timestamp}.xlsx", _xlsx_bytes(sheets), XLSX_CONTENT_TYPE
    return f"{prefix}-{timestamp}.csv", _csv_bytes(sheets), CSV_CONTENT_TYPE
